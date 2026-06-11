"""CSV/Excel import: parsing, column mapping and row grouping (pure functions).

Conversion policy (no guessing, per the compliance ground rules):
- Dates must match the mapping's explicit ``date_format`` (default ISO).
- VAT cells must resolve through the explicit ``vat_value_map`` or be a bare
  category code (S/Z/E/AE). Unknown values are per-row errors, never defaulted.
- Numbers accept thousands separators/whitespace (mechanical cleanup), nothing else.
- Rows sharing an invoice-number cell become the lines of one invoice; their
  header cells (date, buyer, totals…) must agree exactly, else the whole invoice
  is reported inconsistent.
"""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.core.errors import ErrorDetail, build_error_detail
from app.pintae.canonical import VatCategory
from app.pintae.v1.calc import CATEGORY_RATES
from app.schemas.importing import ImportMapping

MAX_ROWS = 5_000

# Codes emitted via build_error_detail here (the i18n completeness test imports this).
IMPORT_ERROR_CODES: frozenset[str] = frozenset(
    {
        "import.file_unsupported",
        "import.file_empty",
        "import.too_many_rows",
        "import.column_missing",
        "import.required_value_missing",
        "import.bad_decimal",
        "import.bad_date",
        "import.vat_value_unknown",
        "import.inconsistent_invoice",
    }
)


class ImportFileError(Exception):
    """File-level problem — nothing row-wise can be reported."""

    def __init__(self, errors: list[ErrorDetail]) -> None:
        super().__init__(errors[0].code if errors else "import file error")
        self.errors = errors


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, Any]]  # header -> raw cell value (str / number / date / None)


@dataclass
class LineData:
    """One file row mapped to canonical line fields (plus its source row number)."""

    row: int
    values: dict[str, Any]


@dataclass
class InvoiceGroup:
    invoice_number: str
    rows: list[int] = field(default_factory=list)
    header: dict[str, Any] = field(default_factory=dict)
    raw_header_cells: dict[str, Any] = field(default_factory=dict)
    lines: list[LineData] = field(default_factory=list)
    errors: list[ErrorDetail] = field(default_factory=list)


def _err(code: str, *, field_name: str | None = None, **params: str) -> ErrorDetail:
    return build_error_detail(code, field=field_name, params=params)


# ------------------------------------------------------------------------ parsing


def parse_file(filename: str, content: bytes) -> ParsedFile:
    name = filename.lower()
    if name.endswith(".csv"):
        parsed = _parse_csv(content)
    elif name.endswith(".xlsx"):
        parsed = _parse_xlsx(content)
    else:
        raise ImportFileError([_err("import.file_unsupported", filename=filename)])

    if not parsed.rows:
        raise ImportFileError([_err("import.file_empty")])
    if len(parsed.rows) > MAX_ROWS:
        raise ImportFileError([_err("import.too_many_rows", limit=str(MAX_ROWS))])
    return parsed


def _parse_csv(content: bytes) -> ParsedFile:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows_iter = (row for row in reader if any(cell.strip() for cell in row))
    try:
        headers = [h.strip() for h in next(rows_iter)]
    except StopIteration:
        raise ImportFileError([_err("import.file_empty")]) from None
    rows = [
        {headers[i]: (row[i].strip() if i < len(row) else "") for i in range(len(headers))}
        for row in rows_iter
    ]
    return ParsedFile(headers=headers, rows=rows)


def _parse_xlsx(content: bytes) -> ParsedFile:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    try:
        first = next(iterator)
    except StopIteration:
        raise ImportFileError([_err("import.file_empty")]) from None
    headers = [str(h).strip() if h is not None else "" for h in first]
    rows = [
        dict(zip(headers, row, strict=False))
        for row in iterator
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    workbook.close()
    return ParsedFile(headers=headers, rows=rows)


# --------------------------------------------------------------------- conversion


def to_decimal(raw: Any) -> Decimal:
    if isinstance(raw, Decimal):
        return raw
    if isinstance(raw, int):
        return Decimal(raw)
    if isinstance(raw, float):
        return Decimal(str(raw))  # shortest-repr round trip; never binary expansion
    text = str(raw).strip().replace(",", "").replace(" ", "")
    if not text:
        raise ValueError("empty")
    return Decimal(text)


def to_date(raw: Any, fmt: str) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    return datetime.strptime(str(raw).strip(), fmt).date()


def to_vat_category(raw: Any, value_map: dict[str, VatCategory]) -> VatCategory:
    text = str(raw).strip()
    if not text:
        raise ValueError("empty")
    normalized = text.lower()
    if normalized in value_map:
        return value_map[normalized]
    try:
        return VatCategory(text.upper())
    except ValueError:
        raise ValueError(text) from None


# ----------------------------------------------------------------------- mapping


_HEADER_FIELDS = (
    "issue_date",
    "due_date",
    "buyer_name",
    "buyer_trn",
    "note",
    "total_excluding_vat",
    "total_vat",
    "total_including_vat",
)
_LINE_FIELDS = (
    "description",
    "quantity",
    "unit_price",
    "discount_amount",
    "vat_treatment",
    "exemption_reason",
    "line_net",
    "line_vat",
)


def check_columns(parsed: ParsedFile, mapping: ImportMapping) -> list[ErrorDetail]:
    mapped = mapping.columns.model_dump(exclude_none=True)
    return [
        _err("import.column_missing", field_name=field_name, column=column)
        for field_name, column in mapped.items()
        if column not in parsed.headers
    ]


def group_rows(
    parsed: ParsedFile, mapping: ImportMapping
) -> tuple[list[InvoiceGroup], list[tuple[int, list[ErrorDetail]]]]:
    """Map+convert each row, group rows into invoices by the invoice-number cell.

    Returns (groups in file order, row-level issues for rows that could not be
    attached to any invoice).
    """
    columns = mapping.columns
    groups: dict[str, InvoiceGroup] = {}
    orphan_issues: list[tuple[int, list[ErrorDetail]]] = []

    for index, raw_row in enumerate(parsed.rows):
        row_no = index + 1  # first DATA row is row 1 from the customer's viewpoint
        number_cell = raw_row.get(columns.invoice_number)
        number = str(number_cell).strip() if number_cell is not None else ""
        if not number:
            orphan_issues.append(
                (
                    row_no,
                    [
                        _err(
                            "import.required_value_missing",
                            field_name="invoice_number",
                            column=columns.invoice_number,
                        )
                    ],
                )
            )
            continue

        group = groups.get(number)
        if group is None:
            group = groups[number] = InvoiceGroup(invoice_number=number)
            _extract_header(group, raw_row, mapping)
        else:
            _check_header_consistency(group, raw_row, mapping)

        group.rows.append(row_no)
        line_values, line_errors = _extract_line(raw_row, row_no, mapping)
        group.lines.append(LineData(row=row_no, values=line_values))
        group.errors.extend(line_errors)

    return list(groups.values()), orphan_issues


def _cell(raw_row: dict[str, Any], column: str | None) -> Any:
    if column is None:
        return None
    return raw_row.get(column)


def _is_empty(value: Any) -> bool:
    return value is None or not str(value).strip()


def _extract_header(group: InvoiceGroup, raw_row: dict[str, Any], mapping: ImportMapping) -> None:
    columns = mapping.columns
    header: dict[str, Any] = {}

    for field_name in _HEADER_FIELDS:
        column = getattr(columns, field_name)
        group.raw_header_cells[field_name] = _cell(raw_row, column) if column else None

    issue_raw = _cell(raw_row, columns.issue_date)
    if _is_empty(issue_raw):
        group.errors.append(
            _err(
                "import.required_value_missing",
                field_name="issue_date",
                column=columns.issue_date,
            )
        )
    else:
        try:
            header["issue_date"] = to_date(issue_raw, mapping.date_format)
        except ValueError:
            group.errors.append(
                _err(
                    "import.bad_date",
                    field_name="issue_date",
                    value=str(issue_raw),
                    expected_format=mapping.date_format,
                )
            )

    if columns.due_date and not _is_empty(_cell(raw_row, columns.due_date)):
        try:
            header["due_date"] = to_date(_cell(raw_row, columns.due_date), mapping.date_format)
        except ValueError:
            group.errors.append(
                _err(
                    "import.bad_date",
                    field_name="due_date",
                    value=str(_cell(raw_row, columns.due_date)),
                    expected_format=mapping.date_format,
                )
            )

    buyer_name = _cell(raw_row, columns.buyer_name)
    header["buyer_name"] = "" if _is_empty(buyer_name) else str(buyer_name).strip()
    if columns.buyer_trn:
        trn = _cell(raw_row, columns.buyer_trn)
        header["buyer_trn"] = None if _is_empty(trn) else str(trn).strip()
    if columns.note and not _is_empty(_cell(raw_row, columns.note)):
        header["note"] = str(_cell(raw_row, columns.note)).strip()

    for field_name in ("total_excluding_vat", "total_vat", "total_including_vat"):
        column = getattr(columns, field_name)
        if column and not _is_empty(_cell(raw_row, column)):
            try:
                header[field_name] = to_decimal(_cell(raw_row, column))
            except (ValueError, InvalidOperation):
                group.errors.append(
                    _err(
                        "import.bad_decimal",
                        field_name=field_name,
                        value=str(_cell(raw_row, column)),
                    )
                )

    group.header = header


def _check_header_consistency(
    group: InvoiceGroup, raw_row: dict[str, Any], mapping: ImportMapping
) -> None:
    """All rows of one invoice must carry identical header cells."""
    columns = mapping.columns
    for field_name in _HEADER_FIELDS:
        column = getattr(columns, field_name)
        if not column:
            continue
        current = _cell(raw_row, column)
        first = group.raw_header_cells.get(field_name)
        current_text = "" if _is_empty(current) else str(current).strip()
        first_text = "" if _is_empty(first) else str(first).strip()
        if current_text != first_text:
            already = any(
                e.code == "import.inconsistent_invoice" and e.field == field_name
                for e in group.errors
            )
            if not already:
                group.errors.append(
                    _err("import.inconsistent_invoice", field_name=field_name, column=column)
                )


def _extract_line(
    raw_row: dict[str, Any], row_no: int, mapping: ImportMapping
) -> tuple[dict[str, Any], list[ErrorDetail]]:
    columns = mapping.columns
    values: dict[str, Any] = {}
    errors: list[ErrorDetail] = []

    description = _cell(raw_row, columns.description)
    values["description"] = "" if _is_empty(description) else str(description).strip()

    for field_name, column in (
        ("quantity", columns.quantity),
        ("unit_price", columns.unit_price),
        ("discount_amount", columns.discount_amount),
        ("declared_net_amount", columns.line_net),
        ("declared_vat_amount", columns.line_vat),
    ):
        if column is None:
            continue
        raw = _cell(raw_row, column)
        if _is_empty(raw):
            if field_name in ("quantity", "unit_price"):
                errors.append(
                    _err(
                        "import.required_value_missing",
                        field_name=f"row[{row_no}].{field_name}",
                        column=column,
                    )
                )
            continue
        try:
            values[field_name] = to_decimal(raw)
        except (ValueError, InvalidOperation):
            errors.append(
                _err(
                    "import.bad_decimal",
                    field_name=f"row[{row_no}].{field_name}",
                    value=str(raw),
                )
            )

    vat_raw = _cell(raw_row, columns.vat_treatment)
    if _is_empty(vat_raw):
        errors.append(
            _err(
                "import.required_value_missing",
                field_name=f"row[{row_no}].vat_treatment",
                column=columns.vat_treatment,
            )
        )
    else:
        try:
            category = to_vat_category(vat_raw, mapping.vat_value_map)
            values["vat_category"] = category
            values["vat_rate"] = CATEGORY_RATES[category]
        except ValueError:
            errors.append(
                _err(
                    "import.vat_value_unknown",
                    field_name=f"row[{row_no}].vat_treatment",
                    value=str(vat_raw).strip(),
                )
            )

    if columns.exemption_reason and not _is_empty(_cell(raw_row, columns.exemption_reason)):
        values["exemption_reason"] = str(_cell(raw_row, columns.exemption_reason)).strip()

    return values, errors
